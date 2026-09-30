"""
中国法律法规 · DeepSeek RAG 问答助手
- 自动读取 data/ 目录下的多部法律文件
- BM25 检索相关法条
- 调用 DeepSeek API 生成回答
- 严格依据知识库内法律，拒绝外部法律
"""
import os
import re
import json
from pathlib import Path
from typing import List, Dict, Any, Optional
from contextlib import asynccontextmanager

import jieba
from rank_bm25 import BM25Okapi
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel
from openai import OpenAI

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# ==================== 配置 ====================
DATA_DIR = os.getenv("DATA_DIR", "data")
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY", "")
DEEPSEEK_BASE_URL = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
DEEPSEEK_MODEL = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")

TOP_K = 10          # 多部法律，召回更多片段
MIN_SCORE = 0.05    # 阈值降低，避免漏召回


# ==================== 文件解析 ====================
def read_text_auto(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030", "utf-16", "big5"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


READERS = {
    ".txt": read_text_auto,
    ".md": read_text_auto,
    ".text": read_text_auto,
    ".pdf": read_pdf,
    ".docx": read_docx,
}


def extract_law_name(filename: str) -> str:
    """从 '06_中华人民共和国道路交通安全法.txt' 提取 '中华人民共和国道路交通安全法'"""
    name = filename
    for ext in (".txt", ".md", ".text", ".pdf", ".docx"):
        if name.lower().endswith(ext):
            name = name[: -len(ext)]
            break
    # 去掉形如 "01_"、"02-"、"03 " 的数字前缀
    name = re.sub(r"^\d+\s*[_\-\s.、]+", "", name).strip()
    return name or filename


def load_documents(data_dir: str) -> List[Dict[str, str]]:
    root = Path(data_dir)
    if not root.exists():
        print(f"[警告] 目录不存在：{data_dir}")
        return []

    docs: List[Dict[str, str]] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        ext = path.suffix.lower()
        if ext not in READERS:
            continue
        try:
            text = READERS[ext](path)
            if text and text.strip():
                law_name = extract_law_name(path.name)
                docs.append({
                    "path": str(path),
                    "filename": path.name,
                    "law_name": law_name,
                    "text": text,
                })
                print(f"[读取] {path.name}  →  法律名：{law_name}  ({len(text)} 字)")
        except Exception as e:
            print(f"[跳过] {path.name}: {e}")
    return docs


# ==================== 文本切分 ====================
# 支持 "第X条"、"第X条之一"、"第X条之二"……
ARTICLE_PATTERN = re.compile(
    r"(第[一二三四五六七八九十百千万零两〇0-9]+条(?:之[一二三四五六七八九十]+)?)"
)


def split_document(doc: Dict[str, str]) -> List[Dict[str, Any]]:
    text = doc["text"].strip()
    law_name = doc["law_name"]
    filename = doc["filename"]
    chunks: List[Dict[str, Any]] = []

    parts = ARTICLE_PATTERN.split(text)

    if len(parts) >= 3:
        preamble = parts[0].strip()
        if len(preamble) > 30:
            chunks.append({
                "law_name": law_name,
                "filename": filename,
                "article": "前言",
                "text": preamble[:800],
            })

        for i in range(1, len(parts), 2):
            art = parts[i].strip()
            body = parts[i + 1].strip() if i + 1 < len(parts) else ""
            full = (art + " " + body).strip()
            if len(full) < 5:
                continue
            # 单条过长时切分（保持原文完整，避免丢失）
            if len(full) > 2000:
                for j in range(0, len(full), 1500):
                    chunks.append({
                        "law_name": law_name,
                        "filename": filename,
                        "article": f"{art}({j // 1500 + 1})",
                        "text": full[j:j + 1500],
                    })
            else:
                chunks.append({
                    "law_name": law_name,
                    "filename": filename,
                    "article": art,
                    "text": full,
                })
    else:
        paragraphs = re.split(r"\n\s*\n", text)
        if len(paragraphs) < 3:
            paragraphs = text.split("\n")
        buf, idx = "", 0
        for p in paragraphs:
            p = p.strip()
            if not p:
                continue
            if len(buf) + len(p) < 600:
                buf += ("\n" if buf else "") + p
            else:
                if buf:
                    idx += 1
                    chunks.append({
                        "law_name": law_name,
                        "filename": filename,
                        "article": f"片段{idx}",
                        "text": buf,
                    })
                buf = p
        if buf:
            idx += 1
            chunks.append({
                "law_name": law_name,
                "filename": filename,
                "article": f"片段{idx}",
                "text": buf,
            })
    return chunks


# ==================== 检索器 ====================
class Retriever:
    def __init__(self, chunks: List[Dict[str, Any]]):
        self.chunks = chunks
        print("[索引] 正在分词...")
        self.corpus_tokens = [
            [t for t in jieba.cut(c["text"]) if t.strip()] for c in chunks
        ]
        self.bm25 = BM25Okapi(self.corpus_tokens) if self.corpus_tokens else None
        print(f"[索引] 完成，共 {len(chunks)} 个片段")

    def search(self, query: str, top_k: int = TOP_K) -> List[Dict[str, Any]]:
        if not self.bm25:
            return []
        q_tokens = [t for t in jieba.cut(query) if t.strip()]
        if not q_tokens:
            return []
        scores = self.bm25.get_scores(q_tokens)
        ranked = sorted(enumerate(scores), key=lambda x: -x[1])[:top_k]
        out = []
        for idx, sc in ranked:
            if sc <= 0:
                continue
            c = dict(self.chunks[idx])
            c["score"] = float(sc)
            out.append(c)
        return out


# ==================== 全局状态 ====================
retriever: Optional[Retriever] = None
client: Optional[OpenAI] = None
chunks_all: List[Dict[str, Any]] = []
law_names_all: List[str] = []
SYSTEM_PROMPT: str = ""


def build_system_prompt(law_names: List[str]) -> str:
    law_list = "\n".join(f"- {n}" for n in sorted(set(law_names)))
    return f"""你是一个严谨的中国法律问答助手。你的知识库只包含以下法律法规的完整条文：

{law_list}

【严格回答规则】
1. 只能依据下面提供的「法条原文」回答问题，不得引用任何未在知识库中的法律、法规、司法解释或案例。
2. 每条结论必须准确标注「法律名称 + 条号」，格式如：
   【中华人民共和国道路交通安全法 第九十一条】
   【中华人民共和国刑法 第一百三十三条之一】
3. 如果提供的法条不足以回答问题，直接回复："根据现有法条，无法回答该问题。"
4. 如果用户的问题与知识库内法律完全无关（如民法、婚姻、劳动、公司等），必须回复：
   "抱歉，我的知识库中不包含相关法律，无法回答该问题。"
5. 使用简洁规范的中文，分点回答，避免客套与废话。
6. 当涉及多部法律时，请分别列明每部法律的对应条文依据。
7. 不得编造条号或内容，所有回答必须能在下面提供的法条原文中找到依据。"""


@asynccontextmanager
async def lifespan(app: FastAPI):
    global retriever, client, chunks_all, law_names_all, SYSTEM_PROMPT

    jieba.initialize()
    print("=" * 64)
    print("启动中...")
    docs = load_documents(DATA_DIR)

    if not docs:
        print(f"[错误] {DATA_DIR}/ 下未找到任何法律文件")
    else:
        for d in docs:
            chunks_all.extend(split_document(d))
            law_names_all.append(d["law_name"])
        retriever = Retriever(chunks_all)
        SYSTEM_PROMPT = build_system_prompt(law_names_all)
        print(f"[法律] 共加载 {len(set(law_names_all))} 部法律：")
        for n in sorted(set(law_names_all)):
            print(f"       · {n}")

    if DEEPSEEK_API_KEY:
        client = OpenAI(api_key=DEEPSEEK_API_KEY, base_url=DEEPSEEK_BASE_URL)
        print(f"[API] DeepSeek 客户端已就绪 · 模型={DEEPSEEK_MODEL}")
    else:
        print("[警告] 未设置 DEEPSEEK_API_KEY，请在 .env 或环境变量中配置")
    print("=" * 64)
    yield
    print("[退出]")


app = FastAPI(title="中国法律 RAG 助手", lifespan=lifespan)


def build_user_prompt(question: str, contexts: List[Dict[str, Any]]) -> str:
    if contexts:
        blocks = []
        for c in contexts:
            blocks.append(
                f"【{c['law_name']} {c['article']}】（来源文件：{c['filename']}）\n{c['text']}"
            )
        ctx = "\n\n".join(blocks)
    else:
        ctx = "（未检索到相关法条）"

    return f"""以下是从知识库中检索到的法条原文：

{ctx}

用户问题：{question}

请严格依据上述法条回答，并标注「法律名称 + 条号」。"""


# ==================== API ====================
class ChatRequest(BaseModel):
    question: str


def ndjson(obj: Dict[str, Any]) -> str:
    return json.dumps(obj, ensure_ascii=False) + "\n"


@app.post("/api/chat")
async def chat(req: ChatRequest):
    q = (req.question or "").strip()
    if not q:
        raise HTTPException(400, "问题不能为空")
    if retriever is None:
        raise HTTPException(500, "知识库未加载，请检查 data/ 目录")
    if client is None:
        raise HTTPException(500, "未配置 DEEPSEEK_API_KEY")

    results = retriever.search(q, TOP_K)
    results = [r for r in results if r["score"] >= MIN_SCORE]
    user_prompt = build_user_prompt(q, results)

    def gen():
        yield ndjson({
            "type": "evidence",
            "data": [
                {
                    "law_name": r["law_name"],
                    "article": r["article"],
                    "filename": r["filename"],
                    "text": r["text"],
                    "score": r["score"],
                }
                for r in results
            ],
        })

        try:
            stream = client.chat.completions.create(
                model=DEEPSEEK_MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_prompt},
                ],
                stream=True,
                temperature=0.1,
            )
            for ev in stream:
                if not ev.choices:
                    continue
                delta = ev.choices[0].delta
                if delta and delta.content:
                    yield ndjson({"type": "token", "data": delta.content})
        except Exception as e:
            yield ndjson({"type": "error", "data": f"调用 DeepSeek 失败：{e}"})

        yield ndjson({"type": "done"})

    return StreamingResponse(gen(), media_type="application/x-ndjson")


@app.get("/api/stats")
async def stats():
    """返回知识库统计信息，方便前端展示"""
    return {
        "chunks": len(chunks_all),
        "laws": sorted(set(law_names_all)),
        "law_count": len(set(law_names_all)),
    }


# ==================== 前端页面 ====================
INDEX_HTML = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>中国法律 · RAG 助手</title>
<style>
:root{--brand:#1d4ed8;--brand2:#2563eb;--ink:#1f2937;--muted:#6b7280;
      --line:#e5e7eb;--bg:#f5f7fb;--card:#fff;--danger:#b91c1c;}
*{box-sizing:border-box;}
body{margin:0;background:var(--bg);color:var(--ink);
     font-family:"PingFang SC","Microsoft YaHei","Helvetica Neue",Arial,sans-serif;}
.wrap{max-width:1380px;margin:0 auto;padding:16px 16px 36px;}
header.hero{background:linear-gradient(135deg,#1e3a8a,#2563eb 60%,#3b82f6);
  color:#fff;border-radius:16px;padding:18px 24px;box-shadow:0 10px 28px rgba(30,58,138,.22);}
header.hero h1{margin:0;font-size:21px;}
header.hero p{margin:6px 0 0;font-size:13px;opacity:.9;}
.badges{margin-top:10px;display:flex;gap:8px;flex-wrap:wrap;}
.badge{font-size:12px;padding:3px 10px;border-radius:999px;
  background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.25);}
.grid{display:grid;grid-template-columns:1.55fr 1fr;gap:16px;margin-top:16px;}
@media(max-width:1000px){.grid{grid-template-columns:1fr;}}
.panel{background:var(--card);border:1px solid var(--line);border-radius:16px;
  box-shadow:0 4px 16px rgba(17,24,39,.05);display:flex;flex-direction:column;overflow:hidden;}
.panel-head{padding:12px 18px;border-bottom:1px solid var(--line);
  font-size:14px;font-weight:600;display:flex;justify-content:space-between;
  align-items:center;background:#fbfcfe;gap:10px;}
.panel-head .sub{font-size:12px;font-weight:400;color:var(--muted);}
#chat{height:560px;overflow-y:auto;padding:18px;background:#fff;}
#chat::-webkit-scrollbar{width:8px;}
#chat::-webkit-scrollbar-thumb{background:#dbe1ea;border-radius:8px;}
.msg{display:flex;gap:10px;margin-bottom:16px;align-items:flex-start;}
.msg .avatar{flex:0 0 34px;height:34px;width:34px;border-radius:10px;
  display:flex;align-items:center;justify-content:center;
  font-size:15px;font-weight:700;color:#fff;}
.msg.user{flex-direction:row-reverse;}
.msg.user .avatar{background:#0ea5e9;}
.msg.bot .avatar{background:linear-gradient(135deg,#1e3a8a,#3b82f6);}
.bubble{max-width:80%;padding:11px 14px;border-radius:12px;font-size:14px;
  line-height:1.75;word-break:break-word;overflow-wrap:anywhere;white-space:pre-wrap;}
.msg.bot .bubble{background:#f3f6fc;border:1px solid #e4ebf7;border-top-left-radius:3px;}
.msg.user .bubble{background:var(--brand2);color:#fff;border-top-right-radius:3px;}
.composer{border-top:1px solid var(--line);padding:12px 14px;background:#fbfcfe;
  display:flex;gap:10px;align-items:flex-end;}
#q{flex:1;border:1px solid #d8dfe9;border-radius:11px;padding:11px 13px;
  font-size:14px;outline:none;font-family:inherit;resize:none;
  background:#fff;max-height:130px;line-height:1.5;transition:.18s;}
#q:focus{border-color:var(--brand2);box-shadow:0 0 0 3px rgba(37,99,235,.12);}
button{font-family:inherit;cursor:pointer;border:none;border-radius:11px;font-size:14px;transition:.16s;}
#send{background:linear-gradient(135deg,#1e3a8a,#2563eb);color:#fff;padding:11px 22px;font-weight:600;}
#send:hover:not(:disabled){filter:brightness(1.08);transform:translateY(-1px);}
#send:disabled{opacity:.55;cursor:not-allowed;}
#stop{background:#fff;color:#b91c1c;border:1px solid #fecaca;padding:11px 16px;font-weight:600;display:none;}
#stop:hover{background:#fef2f2;}
#clear{background:#fff;color:var(--muted);border:1px solid #d8dfe9;padding:6px 13px;font-size:12px;}
#clear:hover{background:#f1f5fb;color:var(--ink);}
.examples{padding:0 14px 14px;background:#fbfcfe;display:flex;flex-wrap:wrap;gap:8px;}
.chip{background:#fff;border:1px solid #dbe3ef;color:#37517e;padding:6px 12px;border-radius:999px;font-size:12.5px;}
.chip:hover{background:#eaf1ff;border-color:#9cb8e8;color:var(--brand);}
.stat{display:flex;gap:16px;padding:9px 18px;border-bottom:1px solid var(--line);
  font-size:12px;color:var(--muted);background:#fbfcfe;flex-wrap:wrap;}
.stat b{color:var(--brand);font-size:13px;}
#evidence{padding:14px;overflow-y:auto;height:665px;background:#fff;}
.ev-card{border:1px solid var(--line);border-left:3px solid var(--brand2);
  border-radius:10px;padding:11px 13px;margin-bottom:11px;background:#fcfdff;animation:fade .28s;}
@keyframes fade{from{opacity:0;transform:translateY(4px);}to{opacity:1;transform:none;}}
.ev-head{display:flex;justify-content:space-between;align-items:flex-start;
  font-weight:700;color:#1e3a8a;font-size:13.5px;gap:8px;}
.ev-head .title{flex:1;line-height:1.4;}
.ev-head .law{color:#7c3aed;font-size:12px;font-weight:500;}
.ev-score{font-size:11.5px;color:#8b95a5;font-weight:500;background:#f0f3f8;
  padding:2px 7px;border-radius:6px;white-space:nowrap;}
.ev-src{font-size:11px;color:#9aa4b2;margin:4px 0 7px;}
.ev-bar-wrap{height:4px;background:#eef1f6;border-radius:3px;overflow:hidden;margin-bottom:9px;}
.ev-bar{height:100%;background:linear-gradient(90deg,#60a5fa,#1d4ed8);border-radius:3px;}
.ev-text{font-size:12.8px;color:#3d4757;line-height:1.72;}
.ev-empty{font-size:13px;color:#98a2b3;text-align:center;padding:26px 14px;
  border:1px dashed var(--line);border-radius:12px;line-height:1.7;}
.typing i{display:inline-block;width:5px;height:5px;margin-right:3px;border-radius:50%;
  background:#8fa8d8;animation:blink 1.2s infinite;}
.typing i:nth-child(2){animation-delay:.2s;}
.typing i:nth-child(3){animation-delay:.4s;}
@keyframes blink{0%,60%,100%{opacity:.25;}30%{opacity:1;}}
</style>
</head>
<body>
<div class="wrap">
  <header class="hero">
    <h1>⚖️ 中国法律 · DeepSeek RAG 问答助手</h1>
    <p>自动加载 data/ 目录下全部法律文件 → BM25 检索 → DeepSeek 流式生成 → 严格依据知识库作答</p>
    <div class="badges">
      <span class="badge" id="badge-laws">📚 知识库加载中...</span>
      <span class="badge">🔍 BM25 本地召回</span>
      <span class="badge">🤖 DeepSeek 流式输出</span>
      <span class="badge">🛡️ 未收录法律拒答</span>
    </div>
  </header>

  <div class="grid">
    <section class="panel">
      <div class="panel-head">
        <span>💬 对话</span>
        <button id="clear">清空对话</button>
      </div>
      <div id="chat"></div>
      <div class="composer">
        <textarea id="q" rows="1" placeholder="请输入法律问题，例如：醉酒驾驶机动车怎么处罚？"></textarea>
        <button id="stop">停止</button>
        <button id="send">发送</button>
      </div>
      <div class="examples" id="examples"></div>
    </section>

    <section class="panel">
      <div class="panel-head">
        <span>📖 RAG 召回的条文</span>
        <span class="sub">检索溯源</span>
      </div>
      <div class="stat">
        <span>召回数：<b id="stat-hit">0</b></span>
        <span>Top 相关度：<b id="stat-score">—</b></span>
      </div>
      <div id="evidence">
        <div class="ev-empty">尚未发起提问<br>召回的法条将在此处可视化展示</div>
      </div>
    </section>
  </div>
</div>

<script>
const chatEl  = document.getElementById('chat');
const inputEl = document.getElementById('q');
const sendBtn = document.getElementById('send');
const stopBtn = document.getElementById('stop');
const clearBtn= document.getElementById('clear');
const evidEl  = document.getElementById('evidence');
const statHit = document.getElementById('stat-hit');
const statScore = document.getElementById('stat-score');
const exWrap  = document.getElementById('examples');
const badgeLaws = document.getElementById('badge-laws');

let sending = false;
let abortController = null;

function escapeHtml(s){
  return String(s).replace(/[&<>"']/g, c => (
    {'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]
  ));
}

function appendMsg(role, text){
  const div = document.createElement('div');
  div.className = 'msg ' + role;
  div.innerHTML = `<div class="avatar">${role==='user'?'我':'法'}</div>
                   <div class="bubble">${text||''}</div>`;
  chatEl.appendChild(div);
  chatEl.scrollTop = chatEl.scrollHeight;
  return div;
}

function renderEvidence(list){
  statHit.textContent = list.length;
  if(!list.length){
    statScore.textContent = '—';
    evidEl.innerHTML = '<div class="ev-empty">未检索到相关法条，已提示模型拒绝回答。</div>';
    return;
  }
  const maxS = Math.max(...list.map(x=>x.score), 0.0001);
  statScore.textContent = list[0].score.toFixed(2);
  evidEl.innerHTML = list.map(item=>{
    const pct = Math.max(8, Math.min(100, item.score / maxS * 100));
    const preview = item.text.length > 260 ? item.text.slice(0,260) + '…' : item.text;
    return `<div class="ev-card">
      <div class="ev-head">
        <span class="title">
          <span class="law">${escapeHtml(item.law_name)}</span><br>
          ${escapeHtml(item.article)}
        </span>
        <span class="ev-score">${item.score.toFixed(2)}</span>
      </div>
      <div class="ev-src">📁 ${escapeHtml(item.filename)}</div>
      <div class="ev-bar-wrap"><div class="ev-bar" style="width:${pct.toFixed(0)}%"></div></div>
      <div class="ev-text">${escapeHtml(preview)}</div>
    </div>`;
  }).join('');
}

async function send(){
  const q = inputEl.value.trim();
  if(!q || sending) return;

  appendMsg('user', escapeHtml(q));
  inputEl.value = '';
  inputEl.style.height = 'auto';

  const botEl = appendMsg('bot', '<span class="typing"><i></i><i></i><i></i></span>');
  const bubble = botEl.querySelector('.bubble');

  sending = true;
  sendBtn.disabled = true;
  stopBtn.style.display = 'inline-block';
  abortController = new AbortController();

  let answer = '';
  let first = true;

  try{
    const resp = await fetch('/api/chat', {
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body: JSON.stringify({question: q}),
      signal: abortController.signal,
    });

    if(!resp.ok){
      const txt = await resp.text();
      bubble.textContent = '请求失败：' + txt;
      return;
    }

    const reader = resp.body.getReader();
    const decoder = new TextDecoder();
    let buf = '';

    while(true){
      const {done, value} = await reader.read();
      if(done) break;
      buf += decoder.decode(value, {stream:true});
      const lines = buf.split('\n');
      buf = lines.pop() || '';
      for(const line of lines){
        const s = line.trim();
        if(!s) continue;
        let msg;
        try{ msg = JSON.parse(s); }catch(e){ continue; }
        if(msg.type === 'evidence'){
          renderEvidence(msg.data || []);
        }else if(msg.type === 'token'){
          if(first){ bubble.textContent = ''; first = false; }
          answer += msg.data;
          bubble.textContent = answer;
          chatEl.scrollTop = chatEl.scrollHeight;
        }else if(msg.type === 'error'){
          bubble.textContent = (answer || '') + '\n\n[错误] ' + msg.data;
        }
      }
    }
  }catch(e){
    if(e.name === 'AbortError'){
      bubble.textContent = (answer || '') + '\n\n[已停止]';
    }else{
      bubble.textContent = '网络错误：' + e.message;
    }
  }finally{
    sending = false;
    sendBtn.disabled = false;
    stopBtn.style.display = 'none';
    abortController = null;
  }
}

sendBtn.addEventListener('click', send);
stopBtn.addEventListener('click', ()=>{ if(abortController) abortController.abort(); });
inputEl.addEventListener('keydown', e=>{
  if(e.key === 'Enter' && !e.shiftKey){ e.preventDefault(); send(); }
});
inputEl.addEventListener('input', ()=>{
  inputEl.style.height = 'auto';
  inputEl.style.height = Math.min(inputEl.scrollHeight, 130) + 'px';
});
clearBtn.addEventListener('click', ()=>{
  chatEl.innerHTML = '';
  evidEl.innerHTML = '<div class="ev-empty">尚未发起提问<br>召回的法条将在此处可视化展示</div>';
  statHit.textContent = '0';
  statScore.textContent = '—';
});

const EXAMPLES = [
  '醉酒驾驶机动车怎么处罚？',
  '刑法第133条之一规定了什么？',
  '行人闯红灯会被罚款吗？',
  '盗窃罪的量刑标准是什么？',
  '行政处罚有哪些种类？',
  '个人信息处理者有哪些义务？',
  '治安管理处罚的程序是怎样的？',
  '刑事案件的强制措施有哪些？'
];
EXAMPLES.forEach(t=>{
  const b = document.createElement('button');
  b.className = 'chip';
  b.textContent = t;
  b.onclick = ()=>{ inputEl.value = t; send(); };
  exWrap.appendChild(b);
});

// 启动时读取知识库信息
fetch('/api/stats').then(r=>r.json()).then(d=>{
  badgeLaws.textContent = `📚 已加载 ${d.law_count} 部法律 · ${d.chunks} 条`;
  console.log('知识库包含的法律：', d.laws);
}).catch(()=>{
  badgeLaws.textContent = '📚 知识库加载失败';
});

appendMsg('bot',
  '您好！我是中国法律 RAG 问答助手。\n\n' +
  '我已加载 data/ 目录下的全部法律文件，包括刑法、刑事诉讼法、治安管理处罚法、行政处罚法、' +
  '道路交通安全法、网络安全法、数据安全法、个人信息保护法等。\n\n' +
  '提问时我会：\n' +
  '· 先用 BM25 检索相关法条（右侧展示）\n' +
  '· 再交给 DeepSeek 依据法条作答\n' +
  '· 每条结论都会标注「法律名称 + 条号」\n\n' +
  '试试点击下方示例问题，或直接输入你的问题。'
);
</script>
</body>
</html>
"""


@app.get("/", response_class=HTMLResponse)
async def index():
    return INDEX_HTML


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=False)